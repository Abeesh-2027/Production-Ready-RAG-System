from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from app.errors import DocumentExtractionError, UnsupportedFileTypeError


@dataclass
class Page:
    """A unit of text with an optional page number (1-based).

    PDFs have real pages. DOCX and TXT have no reliable page concept, so their
    page_number is None and they are loaded as a single page.
    """
    text: str
    page_number: int | None = None


@dataclass
class LoadedDocument:
    filename: str
    file_type: str  
    pages: list[Page]

    @property
    def total_characters(self) -> int:
        return sum(len(p.text) for p in self.pages)


def _load_pdf(path: Path) -> list[Page]:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise DocumentExtractionError("PyMuPDF is not installed on the server.") from exc

    try:
        with fitz.open(path) as pdf:
            if pdf.needs_pass:
                raise DocumentExtractionError("This PDF is password-protected and cannot be read.")
            return [Page(text=page.get_text("text"), page_number=i) for i, page in enumerate(pdf, start=1)]
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract text from the PDF: {exc}") from exc


def _load_docx(path: Path) -> list[Page]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentExtractionError("python-docx is not installed on the server.") from exc

    try:
        doc = Document(str(path))
        parts: list[str] = [p.text for p in doc.paragraphs]
        # Tables are not part of doc.paragraphs, so read them explicitly.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [Page(text="\n".join(parts), page_number=None)]
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract text from the DOCX file: {exc}") from exc


def _load_txt(path: Path) -> list[Page]:
    raw = path.read_bytes()
    encodings = ["utf-8-sig", "latin-1"]
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.insert(0, "utf-16")
    for encoding in encodings:
        try:
            return [Page(text=raw.decode(encoding), page_number=None)]
        except UnicodeError:
            continue
    raise DocumentExtractionError("Could not decode the text file.")  # pragma: no cover


LOADERS: dict[str, Callable[[Path], list[Page]]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_txt,
}


def supported_extensions() -> tuple[str, ...]:
    return tuple(LOADERS)


def load_document(path: Path, original_filename: str | None = None) -> LoadedDocument:
    """Detect the file type from its extension and extract its text."""
    ext = path.suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported types: {', '.join(LOADERS)}."
        )
    pages = loader(path)
    return LoadedDocument(
        filename=original_filename or path.name,
        file_type=ext.lstrip("."),
        pages=pages,
    )
