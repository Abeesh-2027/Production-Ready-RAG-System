from pathlib import Path

import pytest

from app.errors import UnsupportedFileTypeError
from app.rag.chunker import chunk_document, split_text
from app.rag.cleaner import clean_text
from app.rag.loader import LoadedDocument, Page, load_document


def test_txt_loader(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("Hello world.\nSecond line.", encoding="utf-8")
    doc = load_document(f, "a.txt")
    assert doc.file_type == "txt" and "Hello world." in doc.pages[0].text


def test_txt_loader_handles_non_utf8(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("caf\xe9".encode("latin-1"))
    assert "caf" in load_document(f).pages[0].text


def test_docx_loader(tmp_path):
    from docx import Document

    d = Document()
    d.add_paragraph("RAG combines retrieval with generation.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Model"
    t.rows[0].cells[1].text = "Embedding"
    f = tmp_path / "a.docx"
    d.save(str(f))
    text = load_document(f).pages[0].text
    assert "retrieval with generation" in text and "Model | Embedding" in text


def test_unsupported_type(tmp_path):
    f = tmp_path / "a.exe"
    f.write_bytes(b"x")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(f)


def test_cleaner_normalises_whitespace_and_keeps_paragraphs():
    raw = "Title\r\n\r\n\r\nThis   is  a   line\nthat wraps here.\n\n\n\nNext paragraph.​"
    out = clean_text(raw)
    assert out == "Title\n\nThis is a line that wraps here.\n\nNext paragraph."


def test_cleaner_joins_hyphenated_words_and_keeps_numbers():
    out = clean_text("The retrie-\nval step costs $4.50 in 2024.")
    assert "retrieval" in out and "$4.50" in out and "2024" in out


def test_cleaner_empty():
    assert clean_text("") == "" and clean_text("  \n \n ") == ""


def _words(n, prefix="w"):
    return " ".join(f"{prefix}{i}" for i in range(n))


def test_chunker_respects_size_and_overlap():
    text = "\n\n".join(_words(100, f"p{i}_") for i in range(10))  # 10 paragraphs x 100 words
    chunks = split_text(text, chunk_size=300, chunk_overlap=50)
    assert len(chunks) > 1
    for c in chunks:
        assert len(c.split()) <= 300 + 50  # size + carried overlap
    # overlap: the tail of chunk N reappears at the start of chunk N+1
    tail = chunks[0].split()[-50:]
    assert chunks[1].split()[:50] == tail


def test_chunker_keeps_paragraphs_whole_when_they_fit():
    text = "Alpha paragraph here.\n\nBeta paragraph here."
    assert split_text(text, 100, 10) == ["Alpha paragraph here.\n\nBeta paragraph here."]


def test_chunker_splits_giant_paragraph_on_sentences():
    para = " ".join(f"Sentence number {i} is here." for i in range(200))  # 1000 words, one paragraph
    chunks = split_text(para, chunk_size=200, chunk_overlap=0)
    assert len(chunks) >= 5 and all(c.rstrip().endswith(".") for c in chunks)


def test_chunker_overlap_larger_than_size_is_safe():
    chunks = split_text(_words(500), chunk_size=100, chunk_overlap=500)
    assert 1 < len(chunks) < 50  # terminates, no infinite loop


def test_chunk_metadata_and_ids():
    doc = LoadedDocument("rag.pdf", "pdf", [Page("x", 1), Page("y", 2)])
    chunks = chunk_document(doc, [(1, "First page text."), (2, "Second page text.")], "doc123", 800, 150)
    assert [c.chunk_id for c in chunks] == ["doc123_chunk_0000", "doc123_chunk_0001"]
    assert chunks[1].page_number == 2 and chunks[1].filename == "rag.pdf"
    meta = chunks[1].metadata()
    assert meta["document_id"] == "doc123" and meta["chunk_index"] == 1 and meta["page_number"] == 2


def test_chunk_metadata_omits_missing_page():
    doc = LoadedDocument("a.txt", "txt", [Page("x")])
    chunk = chunk_document(doc, [(None, "Some text.")], "d", 800, 150)[0]
    assert "page_number" not in chunk.metadata()
